"""Whole-master listings as Excel or CSV — the legacy View Listing reports, for the masters too large
for the browser to hold (the Item master runs to tens of thousands of rows).

Smaller lists are exported by the app from what's already on screen; only these come from here.

This runs on the same server the tills bill against, so it must never hold the request loop: rows
are read as plain values (no model objects), and the file itself is written on a worker thread.
"""
import asyncio
import csv
import io
from datetime import datetime, timezone
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import Party, Product, ProductAlias

Column = tuple[str, Callable[[dict], Any], int]  # title, value getter over a row dict, Excel column width

_CHUNK = 2000


def _yes(value: Any) -> str:
    return "Yes" if value else "No"


_ITEM_FIELDS = (
    "id", "sku", "name", "barcode", "brand", "category", "item_class", "subclass", "department", "manufacturer",
    "unit", "pack_unit", "pack_size", "price", "rpp", "tax_rate", "avg_cost", "disc_percent", "disc_flat",
    "lock_disc", "variant", "origin", "active", "wholesale_price", "reorder_level",
)
_ITEM_COLUMNS: list[Column] = [
    ("Code", lambda r: r["sku"], 14), ("Item name", lambda r: r["name"], 42), ("Barcode", lambda r: r["barcode"], 16),
    ("Brand", lambda r: r["brand"], 18), ("Category", lambda r: r["category"], 18), ("Class", lambda r: r["item_class"], 16),
    ("Sub-class", lambda r: r["subclass"], 18), ("Department", lambda r: r["department"], 14),
    ("Manufacturer", lambda r: r["manufacturer"], 22), ("Unit", lambda r: r["unit"], 8), ("Pack unit", lambda r: r["pack_unit"], 10),
    ("Units per pack", lambda r: r["pack_size"], 10), ("Sale price", lambda r: r["price"], 11), ("Retail price", lambda r: r["rpp"], 11),
    ("Wholesale price", lambda r: r["wholesale_price"], 11),
    ("GST %", lambda r: r["tax_rate"], 7), ("Average cost", lambda r: r["avg_cost"], 12), ("Item disc %", lambda r: r["disc_percent"], 9),
    ("Item flat disc", lambda r: r["disc_flat"], 10), ("Lock discount", lambda r: _yes(r["lock_disc"]), 9),
    ("Variant", lambda r: r["variant"], 10), ("Imported / local", lambda r: (r["origin"] or "").title() or None, 10),
    ("Active", lambda r: _yes(r["active"]), 7), ("Reorder level", lambda r: r["reorder_level"], 9),
    ("Alternate barcodes", lambda r: r.get("alias_codes"), 30),
]

_PARTY_FIELDS = (
    "code", "name", "category", "phone", "telephone", "email", "address", "city", "area", "sub_area", "contact_person",
    "ntn", "cnic", "s_tax_reg_no", "loyalty_no", "tier", "due_days", "credit_allowed", "credit_limit", "credit_balance", "active",
)
_PARTY_COLUMNS: list[Column] = [
    ("Code", lambda r: r["code"], 10), ("Party name", lambda r: r["name"], 32), ("Category", lambda r: r["category"], 14),
    ("Cell no", lambda r: r["phone"], 15), ("Telephone", lambda r: r["telephone"], 15), ("Email", lambda r: r["email"], 24),
    ("Address", lambda r: r["address"], 30), ("City", lambda r: r["city"], 14), ("Area", lambda r: r["area"], 14),
    ("Sub-area", lambda r: r["sub_area"], 14), ("Contact person", lambda r: r["contact_person"], 18), ("NTN", lambda r: r["ntn"], 12),
    ("CNIC", lambda r: r["cnic"], 16), ("S.Tax Reg No", lambda r: r["s_tax_reg_no"], 16), ("Loyalty no", lambda r: r["loyalty_no"], 12),
    ("Price tier", lambda r: (r["tier"] or "").title(), 10), ("Due days", lambda r: r["due_days"], 8),
    ("Credit allowed", lambda r: _yes(r["credit_allowed"]), 9), ("Balance limit", lambda r: r["credit_limit"], 12),
    ("Credit used", lambda r: r["credit_balance"], 12), ("Active", lambda r: _yes(r["active"]), 7),
]


async def _items() -> tuple[str, list[Column], list[dict]]:
    aliases: dict[str, list[str]] = {}
    for product_id, code in await ProductAlias.all().values_list("product_id", "code"):
        aliases.setdefault(product_id, []).append(code)
    # In chunks, handing the loop back between them, so a scan at a till never waits behind the
    # whole master being read at once.
    rows: list[dict] = []
    offset = 0
    while True:
        chunk = await Product.all().order_by("name", "id").offset(offset).limit(_CHUNK).values(*_ITEM_FIELDS)
        for row in chunk:
            row["alias_codes"] = ", ".join(aliases.get(row["id"], [])) or None
        rows.extend(chunk)
        if len(chunk) < _CHUNK:
            break
        offset += _CHUNK
        await asyncio.sleep(0)
    return "Items", _ITEM_COLUMNS, rows


async def _parties() -> tuple[str, list[Column], list[dict]]:
    return "Parties", _PARTY_COLUMNS, await Party.filter(is_walk_in=False).order_by("name").values(*_PARTY_FIELDS)


LISTINGS = {"items": _items, "parties": _parties}


def _plain(value: Any) -> Any:
    """Numbers stay numbers in Excel; everything else becomes text."""
    if value is None or isinstance(value, (int, float, str)):
        return value
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)


def _write_csv(columns: list[Column], rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c[0] for c in columns])
    for row in rows:
        writer.writerow(["" if (v := getter(row)) is None else v for _, getter, _ in columns])
    # The byte-order mark is what makes Excel read Urdu names and the rupee sign correctly.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _write_xlsx(title: str, columns: list[Column], rows: list[dict]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title)
    for index, (_, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    header = []
    for name, _, _ in columns:
        cell = WriteOnlyCell(sheet, value=name)
        cell.font = Font(bold=True)
        header.append(cell)
    sheet.append(header)
    for row in rows:
        sheet.append([_plain(getter(row)) for _, getter, _ in columns])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


async def build(name: str, fmt: str, branch_label: str) -> tuple[bytes, str, str]:
    """(content, media type, file name)."""
    title, columns, rows = await LISTINGS[name]()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    base = f"{title.lower()}-{branch_label}-{stamp}"
    if fmt == "csv":
        content = await asyncio.to_thread(_write_csv, columns, rows)
        return content, "text/csv; charset=utf-8", f"{base}.csv"
    content = await asyncio.to_thread(_write_xlsx, title, columns, rows)
    return content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{base}.xlsx"
